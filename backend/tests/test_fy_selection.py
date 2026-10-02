"""Backend tests for FY-selection feature and non-live-site endpoints."""
import os
import io
import pytest
import requests
from openpyxl import load_workbook
from dotenv import load_dotenv
from pathlib import Path

load_dotenv(Path(__file__).resolve().parents[2] / "frontend" / ".env")
BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    yield s
    # Restore default FY
    s.post(f"{API}/financial-year", json={"value": "2425", "text": "2024 - 2025"}, timeout=15)


def test_config_default(client):
    r = client.get(f"{API}/config", timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert "financial_year" in j and "financial_year_label" in j and "financial_year_value" in j
    assert "crawler" in j
    assert "target_url" in j


def test_set_fy_2324(client):
    r = client.post(f"{API}/financial-year", json={"value": "2324", "text": "2023 - 2024"}, timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert j["ok"] is True
    assert j["financial_year"]["value"] == "2324"
    assert j["financial_year"]["text"] == "2023 - 2024"
    assert j["financial_year"]["label"] == "2023-2024"
    # config reflects
    cfg = client.get(f"{API}/config", timeout=15).json()
    assert cfg["financial_year_value"] == "2324"
    assert cfg["financial_year"] == "2023-2024"
    # summary reflects
    smry = client.get(f"{API}/summary", timeout=15).json()
    assert smry["financial_year"] == "2023-2024"


def test_set_fy_2526(client):
    r = client.post(f"{API}/financial-year", json={"value": "2526", "text": "2025 - 2026"}, timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert j["ok"] is True
    assert j["financial_year"]["label"] == "2025-2026"
    cfg = client.get(f"{API}/config", timeout=15).json()
    assert cfg["financial_year_value"] == "2526"
    assert cfg["financial_year"] == "2025-2026"


def test_export_filename_uses_current_fy(client):
    # Set to 2526 first
    client.post(f"{API}/financial-year", json={"value": "2526", "text": "2025 - 2026"}, timeout=15)
    r = client.get(f"{API}/export?scope=all", timeout=30)
    assert r.status_code == 200
    cd = r.headers.get("Content-Disposition", "")
    assert "2025-2026" in cd, f"Content-Disposition missing FY: {cd}"
    assert cd.endswith('.xlsx"') or ".xlsx" in cd
    # Validate xlsx is readable
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert len(wb.sheetnames) > 0


def test_records_endpoint(client):
    r = client.get(f"{API}/records", timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert isinstance(j, (list, dict))


def test_errors_endpoint(client):
    r = client.get(f"{API}/errors", timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert "errors" in j
    assert isinstance(j["errors"], list)


def test_crawl_status(client):
    r = client.get(f"{API}/crawl/status", timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert j.get("state") == "idle"
    # financial_year present per request
    # (status may put it under meta or directly; check both)
    has_fy = ("financial_year" in j) or ("financial_year" in (j.get("meta") or {}))
    assert has_fy, f"financial_year not in crawl status: {j}"


def test_reset_when_idle(client):
    r = client.post(f"{API}/reset", timeout=15)
    assert r.status_code == 200
    j = r.json()
    assert j.get("ok") is True
