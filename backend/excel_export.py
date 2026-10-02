"""
Excel export using openpyxl.
- Phone numbers & emails stored as TEXT ('@' number format) -> no scientific notation, no lost leading zeroes.
- Blank values stay blank.
Sheets: "GP Contact Data", "Summary", "Errors".
"""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db
from website_adapter import fy_label

TEXT_FMT = "@"
_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=13, color="1E1B4B")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, ncols, row=1):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[row].height = 24


def _autowidth(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _data_sheet(wb, rows):
    ws = wb.active
    ws.title = "GP Contact Data"
    headers = ["Sr No", "District", "Sub Division", "Block", "Gram Panchayat",
               "Financial Year", "GP Email ID", "GIS Mobile No1", "GIS Mobile No2", "Status"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for i, r in enumerate(rows, start=1):
        row_idx = i + 1
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=r.get("district", ""))
        ws.cell(row=row_idx, column=3, value=r.get("subdivision", ""))
        ws.cell(row=row_idx, column=4, value=r.get("block", ""))
        ws.cell(row=row_idx, column=5, value=r.get("gp", ""))
        ws.cell(row=row_idx, column=6, value=r.get("financial_year", fy_label()))
        # Email + mobiles explicitly as TEXT
        for col, key in ((7, "email"), (8, "mob1"), (9, "mob2")):
            cell = ws.cell(row=row_idx, column=col, value=str(r.get(key) or ""))
            cell.number_format = TEXT_FMT
        ws.cell(row=row_idx, column=10, value=r.get("status", ""))
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = _BORDER
    ws.freeze_panes = "A2"
    _autowidth(ws, [8, 20, 20, 22, 26, 14, 30, 16, 16, 12])


def _summary_sheet(wb, meta):
    ws = wb.create_sheet("Summary")
    ws["A1"] = "GP Contact Data — Crawl Summary"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    counts = db.get_counts()
    pairs = [
        ("Financial Year", fy_label()),
        ("Total Districts", meta.get("total_districts", "")),
        ("Total Sub Divisions", meta.get("total_subdivisions", "")),
        ("Total Blocks", meta.get("total_blocks", "")),
        ("Total Gram Panchayats (Discovered)", counts["discovered"]),
        ("Successfully Processed", counts["successful"]),
        ("Partial Records", counts["partial"]),
        ("Failed Records", counts["failed"]),
        ("Duplicate Records (skipped)", meta.get("duplicates", 0)),
        ("Records Missing Email", counts["missing_email"]),
        ("Records Missing GIS Mobile No1", counts["missing_mob1"]),
        ("Records Missing GIS Mobile No2", counts["missing_mob2"]),
        ("Start Time", meta.get("start_time", "")),
        ("End Time", meta.get("end_time", "")),
        ("Total Processing Time", meta.get("elapsed", "")),
    ]
    r = 3
    for k, v in pairs:
        kc = ws.cell(row=r, column=1, value=k); kc.font = Font(bold=True)
        kc.fill = PatternFill("solid", fgColor="EEF2FF"); kc.border = _BORDER
        vc = ws.cell(row=r, column=2, value=v); vc.border = _BORDER
        vc.alignment = Alignment(horizontal="left")
        r += 1
    _autowidth(ws, [38, 34])


def _errors_sheet(wb):
    ws = wb.create_sheet("Errors")
    headers = ["District", "Sub Division", "Block", "Gram Panchayat", "Financial Year",
               "Error Type", "Error Message", "Retry Count", "Timestamp"]
    ws.append(headers)
    _style_header(ws, len(headers))
    for r in db.get_errors():
        ws.append([r.get("district", ""), r.get("subdivision", ""), r.get("block", ""),
                   r.get("gp", ""), r.get("financial_year", ""), r.get("error_type", ""),
                   r.get("error_message", ""), r.get("retry_count", 0), r.get("timestamp", "")])
    ws.freeze_panes = "A2"
    _autowidth(ws, [18, 18, 20, 24, 14, 18, 40, 12, 26])


def build_workbook(scope: str, meta: dict) -> bytes:
    rows = db.all_records_for_export(scope)
    wb = Workbook()
    _data_sheet(wb, rows)
    _summary_sheet(wb, meta)
    _errors_sheet(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
